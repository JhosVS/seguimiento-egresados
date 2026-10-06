# app/admin/egresados.py
import io

from flask import flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import case, func, or_, select
from sqlalchemy.orm import selectinload

from app.admin import admin
from app.admin.dashboard import _filtros_activos, _filtros_de_url, _opciones_filtro
from app.auth.decorators import requiere_rol
from app.extensions import db
from app.models import Curricula, Egresado, FormacionAcademica, ProgramaEstudio, Usuario
from app.utils.helpers import alcance_personal

NIVEL_NUM = {'egresado': 0, 'bachiller': 1, 'titulado': 2, 'maestria': 3, 'doctorado': 4}

# Máximo nivel de cada egresado como número (0 = sin grado), calculado dentro de SQL
NIVEL_SQL = (select(func.coalesce(func.max(case(
                (FormacionAcademica.tipo == 'doctorado', 4),
                (FormacionAcademica.tipo == 'maestria', 3),
                (FormacionAcademica.tipo == 'titulado', 2),
                (FormacionAcademica.tipo == 'bachiller', 1))), 0))
             .where(FormacionAcademica.id_egresado == Egresado.id_egresado)
             .scalar_subquery())


def _limitar(q, personal):
    """Restringe una consulta (que ya une ProgramaEstudio) al alcance del rol."""
    fac_fijo, prog_fijo = alcance_personal(personal)
    if prog_fijo is not None:
        return q.filter(ProgramaEstudio.id_programa == prog_fijo)
    if fac_fijo is not None:
        return q.filter(ProgramaEstudio.id_facultad == fac_fijo)
    return q


def _en_alcance(personal, e):
    """True si el personal puede ver a este egresado."""
    fac_fijo, prog_fijo = alcance_personal(personal)
    programa = e.curricula.programa
    if prog_fijo is not None:
        return programa.id_programa == prog_fijo
    if fac_fijo is not None:
        return programa.id_facultad == fac_fijo
    return True


def _consulta(personal):
    """Egresados dentro del alcance del personal, con los filtros de la URL aplicados."""
    f = request.args
    q = (Egresado.query
         .join(Curricula, Egresado.id_curricula == Curricula.id_curricula)
         .join(ProgramaEstudio, Curricula.id_programa == ProgramaEstudio.id_programa)
         .outerjoin(Usuario, Usuario.id_egresado == Egresado.id_egresado)
         .options(selectinload(Egresado.formaciones)))
    q = _limitar(q, personal)

    id_facultad, id_programa = _filtros_de_url()      # vacíos si el filtro está apagado
    if id_programa:
        q = q.filter(ProgramaEstudio.id_programa == id_programa)
    elif id_facultad:
        q = q.filter(ProgramaEstudio.id_facultad == id_facultad)

    buscar = f.get('buscar', '').strip()
    if buscar:
        like = f'%{buscar}%'
        q = q.filter(or_(
            Egresado.nombres.ilike(like), Egresado.apellido_paterno.ilike(like),
            Egresado.apellido_materno.ilike(like), Egresado.codigo_universitario.ilike(like),
            Egresado.dni.ilike(like), Egresado.telefono.ilike(like),
            Egresado.correo_personal.ilike(like), Usuario.correo.ilike(like)))
    if f.get('genero') in ('M', 'F'):
        q = q.filter(Egresado.genero == f['genero'])
    if f.get('curricula'):
        q = q.filter(Curricula.nombre == f['curricula'])
    if f.get('anio_ingreso', type=int):
        q = q.filter(Egresado.anio_ingreso == f.get('anio_ingreso', type=int))
    if f.get('anio_egreso', type=int):
        q = q.filter(Egresado.anio_egreso == f.get('anio_egreso', type=int))
    if f.get('letra_periodo'):
        q = q.filter(Egresado.periodo_ingreso.ilike(f'%{f["letra_periodo"]}'))

    nivel = f.get('nivel', '')
    if nivel.startswith('solo_') and nivel[5:] in NIVEL_NUM:
        q = q.filter(NIVEL_SQL == NIVEL_NUM[nivel[5:]])
    elif nivel in NIVEL_NUM:
        q = q.filter(NIVEL_SQL >= NIVEL_NUM[nivel])
    return q


@admin.route('/admin/egresados')
def egresados():
    personal = current_user.personal
    filtros = request.args.to_dict()
    filtros.pop('pagina', None)

    pagina = request.args.get('pagina', 1, type=int)
    paginado = (_consulta(personal)
                .order_by(Egresado.apellido_paterno, Egresado.apellido_materno, Egresado.nombres)
                .paginate(page=pagina, per_page=30, error_out=False))

    # Listas de los desplegables (currículas dentro del alcance del rol)
    q_curr = (db.session.query(Curricula.nombre)
              .join(ProgramaEstudio, Curricula.id_programa == ProgramaEstudio.id_programa))
    curriculas = [r[0] for r in _limitar(q_curr, personal).distinct().order_by(Curricula.nombre)]
    anios_ing = sorted(r[0] for r in db.session.query(Egresado.anio_ingreso).distinct() if r[0])
    anios_eg = sorted(r[0] for r in db.session.query(Egresado.anio_egreso).distinct() if r[0])

    id_facultad, id_programa = _filtros_de_url()
    facultades, programas = (_opciones_filtro(personal, id_facultad)
                             if _filtros_activos() else ([], []))

    return render_template('admin/egresados.html',
        egresados=paginado.items, total=paginado.total, paginacion=paginado,
        hay_filtro=bool(filtros), filtros=filtros,
        curriculas=curriculas, anios_ing=anios_ing, anios_eg=anios_eg,
        facultades=facultades, programas=programas,
        id_facultad=id_facultad, id_programa=id_programa)


@admin.route('/admin/egresados/<int:id>')
def ver_egresado(id):
    e = db.get_or_404(Egresado, id)
    if not _en_alcance(current_user.personal, e):
        flash('No tienes permiso para ver este egresado.', 'danger')
        return redirect(url_for('admin.egresados'))
    return render_template('admin/ver_egresado.html', e=e)


@admin.route('/admin/egresados/exportar')
def exportar_egresados():
    egresados = _consulta(current_user.personal).order_by(Egresado.apellido_paterno).all()

    columnas = [
        ('Código universitario', lambda e: e.codigo_universitario or ''),
        ('Apellido paterno',     lambda e: e.apellido_paterno),
        ('Apellido materno',     lambda e: e.apellido_materno or ''),
        ('Nombres',              lambda e: e.nombres),
        ('Correo institucional', lambda e: e.usuario.correo if e.usuario else ''),
        ('Correo personal',      lambda e: e.correo_personal or ''),
        ('Género',               lambda e: {'M': 'Masculino', 'F': 'Femenino'}.get(e.genero, '')),
        ('Currícula',            lambda e: e.curricula.nombre),
        ('Programa',             lambda e: e.programa.nombre),
        ('Facultad',             lambda e: e.facultad.nombre),
        ('Año de ingreso',       lambda e: e.anio_ingreso or ''),
        ('Período de ingreso',   lambda e: e.periodo_ingreso or ''),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Egresados'
    ws.append([c[0] for c in columnas])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='003876', end_color='003876', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    for e in egresados:
        ws.append([c[1](e) for c in columnas])
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 35)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='egresados.xlsx')


@admin.route('/admin/programas-por-facultad')
def programas_por_facultad():
    id_facultad = request.args.get('id_facultad', type=int)
    programas = (ProgramaEstudio.query.filter_by(id_facultad=id_facultad)
                 .order_by(ProgramaEstudio.nombre).all())
    return jsonify([{'id': p.id_programa, 'nombre': p.nombre} for p in programas])


@admin.route('/admin/egresados/<int:id>/pdf')
@requiere_rol('administrador', 'coordinador')
def exportar_egresado_pdf(id):
    from app.utils.cv_pdf import generar_cv
    e = db.get_or_404(Egresado, id)
    if not _en_alcance(current_user.personal, e):
        flash('No tienes permiso para ver este egresado.', 'danger')
        return redirect(url_for('admin.egresados'))

    certificados = [c for c in e.certificados if c.categoria != 'idioma']
    idiomas = [c for c in e.certificados if c.categoria == 'idioma']
    habilidades = [h for h in e.habilidades if h.tipo == 'habilidad']
    intereses = [h for h in e.habilidades if h.tipo == 'interes']

    pdf = generar_cv(e, e.empleos, e.formaciones, certificados, habilidades, intereses, idiomas)
    nombre = f'Perfil_{e.apellido_paterno}_{e.nombres.split()[0]}.pdf'
    return send_file(pdf, mimetype='application/pdf', as_attachment=True, download_name=nombre)