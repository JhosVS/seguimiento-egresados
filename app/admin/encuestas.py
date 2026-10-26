# app/admin/encuestas.py
import io
from datetime import date

from flask import flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import selectinload

from app.admin import admin
from app.auth.decorators import requiere_rol
from app.extensions import db
from app.models import (Curricula, Egresado, Encuesta, EncuestaEnvio, Facultad, Opcion, Pregunta, ProgramaEstudio, Respuesta, RespuestaOpcion)
from app.utils.helpers import alcance_personal
from app.utils.historial import registrar

ROLES_VER = ('administrador', 'coordinador', 'docente_responsable')
TIPOS = ('texto', 'numero', 'si_no', 'opcion_unica', 'opcion_multiple')
TIPOS_CON_OPCIONES = ('opcion_unica', 'opcion_multiple')


def _fecha(valor):
    try:
        return date.fromisoformat(valor)
    except (TypeError, ValueError):
        return None


def _filtros():
    """Facultad y programa pedidos en la URL, ajustados al alcance del rol."""
    id_facultad = request.args.get('id_facultad', type=int)
    id_programa = request.args.get('id_programa', type=int)
    fac_fijo, prog_fijo = alcance_personal(current_user.personal)
    if prog_fijo is not None:          # docente responsable: no filtra, ya está acotado
        return None, None
    if fac_fijo is not None:           # coordinador: siempre dentro de su facultad
        id_facultad = fac_fijo
    return id_facultad, id_programa


def _acotar(q, id_facultad, id_programa):
    """Aplica el alcance del rol y los filtros a una consulta que ya une ProgramaEstudio."""
    fac_fijo, prog_fijo = alcance_personal(current_user.personal)
    if prog_fijo is not None:
        q = q.filter(ProgramaEstudio.id_programa == prog_fijo)
    elif fac_fijo is not None:
        q = q.filter(ProgramaEstudio.id_facultad == fac_fijo)

    if id_programa:
        q = q.filter(ProgramaEstudio.id_programa == id_programa)
    elif id_facultad:
        q = q.filter(ProgramaEstudio.id_facultad == id_facultad)
    return q


def _envios(id_encuesta, id_facultad=None, id_programa=None):
    """Envíos (un egresado = un envío) de la encuesta, dentro del alcance del rol."""
    q = (EncuestaEnvio.query
         .filter(EncuestaEnvio.id_encuesta == id_encuesta)
         .join(Egresado, EncuestaEnvio.id_egresado == Egresado.id_egresado)
         .join(Curricula, Egresado.id_curricula == Curricula.id_curricula)
         .join(ProgramaEstudio, Curricula.id_programa == ProgramaEstudio.id_programa))
    return _acotar(q, id_facultad, id_programa)


def _respuestas(id_encuesta, id_facultad=None, id_programa=None):
    """Respuestas de la encuesta, dentro del alcance del rol y con los filtros aplicados."""
    q = (Respuesta.query
         .join(EncuestaEnvio, Respuesta.id_envio == EncuestaEnvio.id_envio)
         .filter(EncuestaEnvio.id_encuesta == id_encuesta)
         .join(Egresado, EncuestaEnvio.id_egresado == Egresado.id_egresado)
         .join(Curricula, Egresado.id_curricula == Curricula.id_curricula)
         .join(ProgramaEstudio, Curricula.id_programa == ProgramaEstudio.id_programa))
    return _acotar(q, id_facultad, id_programa)


# ------------------------------------------------------------
# Listado y creación
# ------------------------------------------------------------
@admin.route('/admin/encuestas')
@requiere_rol(*ROLES_VER)
def encuestas():
    # Cerrar automáticamente las encuestas cuya fecha fin ya pasó
    vencidas = Encuesta.query.filter(
        Encuesta.estado == 'publicada',
        Encuesta.fecha_fin.isnot(None),
        Encuesta.fecha_fin < date.today()).all()
    for enc in vencidas:
        enc.estado = 'cerrada'
    if vencidas:
        db.session.commit()

    rol_activo = current_user.rol
    q = Encuesta.query
    if rol_activo != 'administrador':
        q = q.filter(Encuesta.estado != 'borrador')
    lista = q.order_by(Encuesta.anio.desc(), Encuesta.id_encuesta.desc()).all()
    return render_template('admin/encuestas.html', encuestas=lista, rol_activo=rol_activo)


@admin.route('/admin/encuestas/nueva', methods=['GET', 'POST'])
@requiere_rol('administrador')
def nueva_encuesta():
    if request.method == 'POST':
        nombre = request.form.get('nombre', '').strip()
        anio = request.form.get('anio', type=int)
        inicio = _fecha(request.form.get('fecha_inicio'))
        fin = _fecha(request.form.get('fecha_fin'))

        if not nombre or not anio:
            flash('Completa el nombre y el año.', 'warning')
        elif inicio and fin and fin < inicio:
            flash('La fecha de fin no puede ser anterior a la de inicio.', 'warning')
        else:
            enc = Encuesta(nombre=nombre,
                           descripcion=request.form.get('descripcion', '').strip() or None,
                           anio=anio, fecha_inicio=inicio, fecha_fin=fin, estado='borrador',
                           id_creado_por=current_user.id_personal)
            db.session.add(enc)

            # Hostorial auditoria
            registrar('creacion', 'encuestas', f'Creó la encuesta "{enc.nombre}"')

            db.session.commit()
            flash('Encuesta creada. Ahora agrega las preguntas.', 'success')
            return redirect(url_for('admin.detalle_encuesta', id=enc.id_encuesta))

    return render_template('admin/nueva_encuesta.html', anio_actual=date.today().year)


# ------------------------------------------------------------
# Detalle y preguntas
# ------------------------------------------------------------
@admin.route('/admin/encuestas/<int:id>')
@requiere_rol(*ROLES_VER)
def detalle_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    rol_activo = current_user.rol

    if enc.estado == 'borrador' and rol_activo != 'administrador':
        flash('Esta encuesta aún no está disponible.', 'warning')
        return redirect(url_for('admin.encuestas'))

    preguntas = Pregunta.query.filter_by(id_encuesta=id).order_by(Pregunta.orden).all()
    total_respuestas = EncuestaEnvio.query.filter_by(id_encuesta=id).count()
    return render_template('admin/detalle_encuesta.html', encuesta=enc, preguntas=preguntas,
                           total_respuestas=total_respuestas, rol_activo=rol_activo)


@admin.route('/admin/encuestas/<int:id>/preguntas/nueva', methods=['POST'])
@requiere_rol('administrador')
def nueva_pregunta(id):
    enc = db.get_or_404(Encuesta, id)
    if enc.estado != 'borrador':
        flash('No se pueden agregar preguntas a una encuesta publicada o cerrada.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))

    tipo = request.form.get('tipo', '').strip()
    texto = request.form.get('texto', '').strip()
    if not texto or tipo not in TIPOS:
        flash('Datos de la pregunta inválidos.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))

    opciones = []
    if tipo in TIPOS_CON_OPCIONES:
        opciones = [l.strip() for l in request.form.get('opciones', '').splitlines() if l.strip()]
        if len(opciones) < 2:
            flash('Las preguntas de opción necesitan al menos 2 opciones válidas.', 'warning')
            return redirect(url_for('admin.detalle_encuesta', id=id))

    max_orden = db.session.query(db.func.max(Pregunta.orden)).filter_by(id_encuesta=id).scalar() or 0
    pregunta = Pregunta(id_encuesta=id, texto=texto, tipo=tipo, orden=max_orden + 1,
                        obligatoria=bool(request.form.get('obligatoria')))
    pregunta.opciones = [Opcion(texto=t, orden=i) for i, t in enumerate(opciones, start=1)]
    db.session.add(pregunta)
    db.session.commit()

    flash('Pregunta agregada.', 'success')
    return redirect(url_for('admin.detalle_encuesta', id=id))


@admin.route('/admin/encuestas/<int:id>/preguntas/<int:id_pregunta>/eliminar', methods=['POST'])
@requiere_rol('administrador')
def eliminar_pregunta(id, id_pregunta):
    enc = db.get_or_404(Encuesta, id)
    if enc.estado != 'borrador':
        flash('No se pueden eliminar preguntas de una encuesta publicada o cerrada.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))

    pregunta = db.get_or_404(Pregunta, id_pregunta)
    if pregunta.id_encuesta != id:
        flash('No corresponde a esta encuesta.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))
    if Respuesta.query.filter_by(id_pregunta=id_pregunta).first():
        flash('No se puede eliminar una pregunta que ya tiene respuestas.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))

    db.session.delete(pregunta)      # sus opciones se borran en cascada
    db.session.commit()
    flash('Pregunta eliminada.', 'info')
    return redirect(url_for('admin.detalle_encuesta', id=id))


# ------------------------------------------------------------
# Estado de la encuesta
# ------------------------------------------------------------
@admin.route('/admin/encuestas/<int:id>/publicar', methods=['POST'])
@requiere_rol('administrador')
def publicar_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    if not Pregunta.query.filter_by(id_encuesta=id).first():
        flash('Agrega al menos una pregunta antes de publicar.', 'warning')
        return redirect(url_for('admin.detalle_encuesta', id=id))
    enc.estado = 'publicada'

    # Historial auditoria
    registrar('actualizacion', 'encuestas', f'Publicó la encuesta "{enc.nombre}"')

    db.session.commit()
    flash('Encuesta publicada. Los egresados ya pueden responderla.', 'success')
    return redirect(url_for('admin.detalle_encuesta', id=id))


@admin.route('/admin/encuestas/<int:id>/cerrar', methods=['POST'])
@requiere_rol('administrador')
def cerrar_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    enc.estado = 'cerrada'

    # Historial auditoria
    registrar('actualizacion', 'encuestas', f'Cerró la encuesta "{enc.nombre}"')

    db.session.commit()
    flash('Encuesta cerrada. Ya no se puede responder.', 'info')
    return redirect(url_for('admin.detalle_encuesta', id=id))


@admin.route('/admin/encuestas/<int:id>/eliminar', methods=['POST'])
@requiere_rol('administrador')
def eliminar_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    if EncuestaEnvio.query.filter_by(id_encuesta=id).first():
        flash('No se puede eliminar una encuesta que ya tiene respuestas.', 'danger')
        return redirect(url_for('admin.detalle_encuesta', id=id))

    # Historial auditoria
    registrar('eliminacion', 'encuestas', f'Eliminó la encuesta "{enc.nombre}"')

    db.session.delete(enc)           # preguntas y opciones se borran en cascada
    db.session.commit()
    flash('Encuesta eliminada.', 'info')
    return redirect(url_for('admin.encuestas'))


# ------------------------------------------------------------
# Resultados
# ------------------------------------------------------------
@admin.route('/admin/encuestas/<int:id>/resultados')
@requiere_rol(*ROLES_VER)
def resultados_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    preguntas = Pregunta.query.filter_by(id_encuesta=id).order_by(Pregunta.orden).all()
    id_facultad, id_programa = _filtros()

    total_respondientes = _envios(id, id_facultad, id_programa).count()

    resultados = []
    for p in preguntas:
        item = {'pregunta': p, 'tipo': p.tipo}
        base = _respuestas(id, id_facultad, id_programa).filter(Respuesta.id_pregunta == p.id_pregunta)

        if p.tipo in TIPOS_CON_OPCIONES:
            conteo = (base.join(RespuestaOpcion, RespuestaOpcion.id_respuesta == Respuesta.id_respuesta)
                      .join(Opcion, RespuestaOpcion.id_opcion == Opcion.id_opcion)
                      .with_entities(Opcion.texto, db.func.count(RespuestaOpcion.id_respuesta))
                      .group_by(Opcion.id_opcion, Opcion.texto, Opcion.orden)
                      .order_by(Opcion.orden).all())
            item['labels'] = [c[0] for c in conteo]
            item['data'] = [c[1] for c in conteo]

        elif p.tipo == 'si_no':
            item['labels'] = ['Sí', 'No']
            item['data'] = [base.filter(Respuesta.respuesta_si_no == v).count() for v in (1, 0)]

        elif p.tipo == 'numero':
            valores = [float(r[0]) for r in base.filter(Respuesta.respuesta_numero.isnot(None))
                       .with_entities(Respuesta.respuesta_numero).all()]
            item['valores'] = valores
            item['promedio'] = round(sum(valores) / len(valores), 2) if valores else None
            item['minimo'] = min(valores) if valores else None
            item['maximo'] = max(valores) if valores else None

        elif p.tipo == 'texto':
            item['respuestas_texto'] = [r[0] for r in base
                .filter(Respuesta.respuesta_texto.isnot(None), Respuesta.respuesta_texto != '')
                .with_entities(Respuesta.respuesta_texto).all()]

        resultados.append(item)

    # Desplegables de filtro, dentro del alcance del rol
    fac_fijo, _ = alcance_personal(current_user.personal)
    q_fac = Facultad.query.order_by(Facultad.nombre)
    if fac_fijo is not None:
        q_fac = q_fac.filter_by(id_facultad=fac_fijo)
    facultades = q_fac.all()
    programas = (ProgramaEstudio.query.filter_by(id_facultad=id_facultad)
                 .order_by(ProgramaEstudio.nombre).all()) if id_facultad else []

    return render_template('admin/resultados_encuesta.html',
        encuesta=enc, resultados=resultados, total_respondientes=total_respondientes,
        facultades=facultades, programas=programas,
        id_facultad=id_facultad, id_programa=id_programa, rol_activo=current_user.rol)


@admin.route('/admin/encuestas/<int:id>/programas-de-facultad')
@requiere_rol('administrador', 'coordinador')
def programas_de_facultad(id):
    id_facultad = request.args.get('id_facultad', type=int)
    fac_fijo, _ = alcance_personal(current_user.personal)
    if fac_fijo is not None and id_facultad != fac_fijo:
        return jsonify([])
    programas = (ProgramaEstudio.query.filter_by(id_facultad=id_facultad)
                 .order_by(ProgramaEstudio.nombre).all())
    return jsonify([{'id': p.id_programa, 'nombre': p.nombre} for p in programas])


# ------------------------------------------------------------
# Exportar a Excel
# ------------------------------------------------------------
@admin.route('/admin/encuestas/<int:id>/exportar')
@requiere_rol(*ROLES_VER)
def exportar_encuesta(id):
    enc = db.get_or_404(Encuesta, id)
    preguntas = Pregunta.query.filter_by(id_encuesta=id).order_by(Pregunta.orden).all()
    id_facultad, id_programa = _filtros()

    filas = (_respuestas(id, id_facultad, id_programa)
             .add_columns(EncuestaEnvio.id_egresado)
             .options(selectinload(Respuesta.opciones).selectinload(RespuestaOpcion.opcion))
             .all())
    mapa = {}
    for resp, id_egresado in filas:
        mapa[(id_egresado, resp.id_pregunta)] = resp
    ids_egresados = sorted({id_egresado for _, id_egresado in filas})

    def formatear(resp, tipo):
        if resp is None:
            return ''
        if tipo == 'texto':
            return resp.respuesta_texto or ''
        if tipo == 'numero':
            return float(resp.respuesta_numero) if resp.respuesta_numero is not None else ''
        if tipo == 'si_no':
            return {1: 'Sí', 0: 'No'}.get(resp.respuesta_si_no, '')
        if tipo in TIPOS_CON_OPCIONES:
            partes = []
            for ro in resp.opciones:
                texto = ro.opcion.texto
                partes.append(f'{texto}: {ro.texto_otro}' if ro.texto_otro else texto)
            return ', '.join(partes)
        return ''

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resultados'
    ws.append(['N°'] + [p.texto for p in preguntas])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='003876', end_color='003876', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    for i, id_egresado in enumerate(ids_egresados, start=1):
        ws.append([i] + [formatear(mapa.get((id_egresado, p.id_pregunta)), p.tipo)
                         for p in preguntas])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 40)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)

    sufijo = ''
    if id_programa:
        p = db.session.get(ProgramaEstudio, id_programa)
        sufijo = f'_{p.codigo}' if p and p.codigo else ''
    elif id_facultad:
        f = db.session.get(Facultad, id_facultad)
        sufijo = f'_{f.codigo}' if f and f.codigo else ''
    nombre = f"encuesta_{enc.anio}_{enc.nombre.replace(' ', '_')[:30]}{sufijo}.xlsx"

    # Historial aduditoria
    registrar('consulta', 'encuestas', f'Exportó a Excel las respuestas de la encuesta "{enc.nombre}"', guardar=True)

    return send_file(mem, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=nombre)